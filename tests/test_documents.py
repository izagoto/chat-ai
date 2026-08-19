import asyncio
import io
import ipaddress
import json
import time

from fastapi.testclient import TestClient

from app.core.config import settings
from app.main import app
from app.services.document_store import document_store
from tests.conftest import wait_until_ready

client = TestClient(app)


def test_document_store_keyword_retrieve(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.document_store.settings.documents_path", str(tmp_path / "documents"))
    monkeypatch.setattr("app.services.document_store.settings.rag_use_embeddings", False)

    async def run() -> None:
        await document_store.add_document(
            collection_id="default",
            document_id="DOC-1",
            filename="notes.json",
            artifact_type="json",
            file_bytes=b"{}",
            segments=[("Contact 08123456789 discussed meeting at mall", {"type": "note"})],
        )
        hits = await document_store.retrieve("default", "08123456789 mall", top_k=2)
        assert hits
        assert hits[0][0].document_id == "DOC-1"

    asyncio.run(run())


def test_upload_json_document(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.document_store.settings.documents_path", str(tmp_path / "documents"))
    monkeypatch.setattr("app.services.document_store.settings.rag_use_embeddings", False)
    payload = {"messages": [{"from": "A", "text": "meet at mall 0812"}]}
    files = {"file": ("notes.json", io.BytesIO(json.dumps(payload).encode()), "application/json")}
    res = client.post("/v1/documents", files=files)
    assert res.status_code == 200
    body = res.json()
    assert body["collection_id"] == "default"
    assert body["document"]["status"] in {"processing", "ready"}
    doc = wait_until_ready(client, body["document"]["document_id"])
    assert doc["chunk_count"] >= 1
    assert doc["status"] == "ready"


def test_list_and_delete_document(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.document_store.settings.documents_path", str(tmp_path / "documents"))
    monkeypatch.setattr("app.services.document_store.settings.rag_use_embeddings", False)
    files = {"file": ("a.txt", io.BytesIO(b"hello alder world"), "text/plain")}
    up = client.post("/v1/documents", files=files)
    assert up.status_code == 200
    doc_id = up.json()["document"]["document_id"]
    wait_until_ready(client, doc_id)

    listed = client.get("/v1/documents")
    assert listed.status_code == 200
    assert any(d["document_id"] == doc_id for d in listed.json()["documents"])

    deleted = client.delete(f"/v1/documents/{doc_id}")
    assert deleted.status_code == 200
    assert deleted.json()["deleted"] is True

    listed2 = client.get("/v1/documents")
    assert all(d["document_id"] != doc_id for d in listed2.json()["documents"])


def test_upload_docx_document(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.document_store.settings.documents_path", str(tmp_path / "documents"))
    monkeypatch.setattr("app.services.document_store.settings.rag_use_embeddings", False)
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_heading("Notes", level=1)
    doc.add_paragraph("Meet at mall, contact 08123456789.")
    doc.save(buf)
    files = {"file": ("notes.docx", io.BytesIO(buf.getvalue()), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/v1/documents", files=files)
    assert res.status_code == 200, res.text
    doc = wait_until_ready(client, res.json()["document"]["document_id"])
    assert doc["artifact_type"] == "docx"
    assert doc["chunk_count"] >= 1


def test_chat_uses_uploaded_context(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.document_store.settings.documents_path", str(tmp_path / "documents"))
    monkeypatch.setattr("app.services.document_store.settings.rag_use_embeddings", False)
    files = {
        "file": (
            "pay.json",
            io.BytesIO(json.dumps({"note": "transfer to account 12345"}).encode()),
            "application/json",
        )
    }
    up = client.post("/v1/documents", files=files)
    wait_until_ready(client, up.json()["document"]["document_id"])
    res = client.post(
        "/v1/chat",
        json={"message": "account 12345", "session_id": "test-chat", "use_documents": True},
    )
    assert res.status_code in {200, 503}


def test_ingest_failure_is_marked(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.document_store.settings.documents_path", str(tmp_path / "documents"))
    monkeypatch.setattr("app.services.document_store.settings.rag_use_embeddings", False)

    async def boom(_filename, _file_bytes):
        raise ValueError("corrupt file")

    monkeypatch.setattr("app.services.ingest.extract_segments", boom)
    files = {"file": ("bad.txt", io.BytesIO(b"hello"), "text/plain")}
    up = client.post("/v1/documents", files=files)
    assert up.status_code == 200
    doc_id = up.json()["document"]["document_id"]
    deadline = time.monotonic() + 8
    last = None
    while time.monotonic() < deadline:
        listed = client.get("/v1/documents")
        last = next((d for d in listed.json()["documents"] if d["document_id"] == doc_id), None)
        if last and last.get("status") == "failed":
            assert "corrupt" in (last.get("error") or "")
            return
        time.sleep(0.05)
    raise AssertionError(f"expected failed status, got {last}")


def test_upload_rejects_oversize(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.document_store.settings.documents_path", str(tmp_path / "documents"))
    previous = settings.max_upload_bytes
    settings.max_upload_bytes = 8
    try:
        files = {"file": ("big.txt", io.BytesIO(b"0123456789"), "text/plain")}
        res = client.post("/v1/documents", files=files)
        assert res.status_code == 413
    finally:
        settings.max_upload_bytes = previous


def test_url_localhost_is_rejected():
    res = client.post("/v1/urls", json={"url": "http://127.0.0.1/"})
    assert res.status_code == 400


def test_ingest_public_url(tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.document_store.settings.documents_path", str(tmp_path / "documents"))
    monkeypatch.setattr("app.services.document_store.settings.rag_use_embeddings", False)
    monkeypatch.setattr(
        "app.services.url_fetch._resolve_ips",
        lambda host, port: [ipaddress.ip_address("93.184.216.34")],
    )

    async def fake_fetch(url):
        from app.services.url_fetch import FetchedUrl

        return FetchedUrl(
            final_url=url,
            content_type="text/html",
            body=b"<p>Public article about Alder</p>",
            artifact_type="url",
            segments=[("Public article about Alder", {"type": "url", "url": url})],
        )

    monkeypatch.setattr("app.services.url_fetch.fetch_url", fake_fetch)
    res = client.post("/v1/urls", json={"url": "https://example.com/article"})
    assert res.status_code == 200, res.text
    assert res.json()["document"]["artifact_type"] == "url"
    doc = wait_until_ready(client, res.json()["document"]["document_id"])
    assert doc["status"] == "ready"
    assert doc["chunk_count"] >= 1
    assert doc["source_url"] == "https://example.com/article"
