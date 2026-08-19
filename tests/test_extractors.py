import json
from pathlib import Path

from app.services.extractors.docx import extract_docx
from app.services.extractors.json_csv import extract_csv, extract_json, extract_text
from app.services.extractors.router import extract_file


def test_extract_json_chat_like(tmp_path: Path):
    path = tmp_path / "chat.json"
    path.write_text(
        json.dumps({"messages": [{"from": "A", "text": "hello"}, {"from": "B", "text": "meet at 9"}]}),
        encoding="utf-8",
    )
    segments = extract_json(path)
    assert segments
    assert "hello" in segments[0].text


def test_extract_csv_rows(tmp_path: Path):
    path = tmp_path / "data.csv"
    path.write_text("name,phone\nAlice,0812\nBob,0813\n", encoding="utf-8")
    segments = extract_csv(path)
    assert segments
    assert "Alice" in segments[0].text
    assert segments[0].location["row_start"] == 1


def test_extract_text_file(tmp_path: Path):
    path = tmp_path / "note.txt"
    path.write_text("Investigator note about suspect contact.", encoding="utf-8")
    segments = extract_text(path)
    assert len(segments) == 1
    assert "Investigator" in segments[0].text


def test_extract_docx_headings_and_table(tmp_path: Path):
    from docx import Document

    path = tmp_path / "brief.docx"
    doc = Document()
    doc.add_heading("Incident summary", level=1)
    doc.add_paragraph("The suspect met a contact at the mall on Friday.")
    doc.add_heading("Phone numbers", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Phone"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "0812"
    doc.save(str(path))

    artifact, segments = extract_file(path)
    assert artifact == "docx"
    texts = "\n".join(seg.text for seg in segments)
    assert "mall" in texts
    assert "Alice" in texts
    assert any(seg.location.get("type") == "docx_section" for seg in segments)
    assert any(seg.location.get("type") == "docx_table" for seg in segments)
    assert extract_docx(path)


def test_legacy_doc_is_rejected(tmp_path: Path):
    path = tmp_path / "old.doc"
    path.write_bytes(b"not a real word file")
    try:
        extract_file(path)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert ".docx" in str(exc)
